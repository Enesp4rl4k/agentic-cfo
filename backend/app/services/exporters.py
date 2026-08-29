"""
Report Exporters — Single Responsibility (SRP) & Open/Closed (OCP) Principles.

Extracts report rendering from agent orchestration:
- ExcelReportExporter: Builds formatted openpyxl workbooks with financial styling.
- DashboardExporter: Builds normalized frontend JSON structures.
"""
from __future__ import annotations

import json
import os
from typing import Any

from app.core.financial import cents_to_amount
from app.core.interfaces import IReportExporter


def _fmt(cents: float) -> str:
    """Format integer cents into a currency string ($XX,XXX.XX)."""
    amt = cents_to_amount(cents)
    return f"${amt:,.2f}"


class ExcelReportExporter(IReportExporter):
    """Generates styled Excel (.xlsx) financial workbooks."""

    format_name: str = "excel"

    def export(
        self,
        pnl: dict[str, Any],
        cashflow: dict[str, Any],
        forecast: dict[str, Any],
        output_path: str,
        **kwargs: Any,
    ) -> str:
        """Render multi-sheet financial report to an .xlsx file."""
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        wb = openpyxl.Workbook()

        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        subheader_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")

        # ── Sheet 1: P&L ──────────────────────────────────────────────────────────
        ws_pnl = wb.active
        ws_pnl.title = "Profit & Loss"
        ws_pnl.append(["Line Item", "Amount ($)"])
        for cell in ws_pnl[1]:
            cell.fill = header_fill
            cell.font = header_font

        pnl_rows = [
            ("Gross Revenue", pnl.get("revenue", 0)),
            ("COGS", -pnl.get("cogs", 0)),
            ("Gross Profit", pnl.get("gross_profit", 0)),
            ("Operating Expenses", -pnl.get("operating_expenses", 0)),
            ("Net Income", pnl.get("net_income", 0)),
            ("Gross Margin (%)", pnl.get("gross_margin_pct", 0.0)),
            ("Net Margin (%)", pnl.get("net_margin_pct", 0.0)),
        ]
        for label, value in pnl_rows:
            if "Margin" in label:
                ws_pnl.append([label, round(value, 2)])
            else:
                ws_pnl.append([label, _fmt(int(value)) if isinstance(value, (int, float)) else value])

        ws_pnl.column_dimensions["A"].width = 35
        ws_pnl.column_dimensions["B"].width = 18

        if pnl.get("narrative"):
            ws_pnl.append([])
            ws_pnl.append(["CFO Commentary"])
            ws_pnl.cell(row=ws_pnl.max_row, column=1).font = Font(bold=True)
            ws_pnl.append([pnl["narrative"]])
            ws_pnl.cell(row=ws_pnl.max_row, column=1).alignment = Alignment(wrap_text=True)
            ws_pnl.row_dimensions[ws_pnl.max_row].height = 80

        # ── Sheet 2: Cash Flow ────────────────────────────────────────────────────
        ws_cf = wb.create_sheet("Cash Flow")
        ws_cf.append(["Activity", "Amount ($)"])
        for cell in ws_cf[1]:
            cell.fill = header_fill
            cell.font = header_font

        cf_rows = [
            ("Operating Cash Flow", cashflow.get("operating", 0)),
            ("  Cash Inflows", cashflow.get("operating_in", 0)),
            ("  Cash Outflows", -cashflow.get("operating_out", 0)),
            ("Investing Cash Flow", cashflow.get("investing", 0)),
            ("Financing Cash Flow", cashflow.get("financing", 0)),
            ("Net Cash Change", cashflow.get("net_change", 0)),
        ]
        for label, value in cf_rows:
            ws_cf.append([label, _fmt(int(value))])

        ws_cf.append([])
        ws_cf.append(["Monthly Cash Flow"])
        ws_cf.cell(row=ws_cf.max_row, column=1).font = Font(bold=True)
        ws_cf.append(["Month", "Cash In ($)", "Cash Out ($)", "Net ($)"])
        for cell in ws_cf[ws_cf.max_row]:
            cell.fill = subheader_fill

        for entry in cashflow.get("monthly_series", []):
            ws_cf.append([
                entry["month"],
                _fmt(entry["in"]),
                _fmt(entry["out"]),
                _fmt(entry["net"]),
            ])

        ws_cf.column_dimensions["A"].width = 30
        for col in ["B", "C", "D"]:
            ws_cf.column_dimensions[col].width = 16

        # ── Sheet 3: Forecast ─────────────────────────────────────────────────────
        ws_fc = wb.create_sheet("Forecast")
        ws_fc.append(["12-Month Financial Forecast"])
        ws_fc.cell(row=1, column=1).font = Font(bold=True, size=13)
        ws_fc.append([])

        for scenario in forecast.get("scenarios", {}).values():
            ws_fc.append([scenario["label"], scenario.get("description", "")])
            ws_fc.cell(row=ws_fc.max_row, column=1).font = Font(bold=True)
            ws_fc.append(["Month", "Cash In ($)", "Cash Out ($)", "Net ($)"])
            for cell in ws_fc[ws_fc.max_row]:
                cell.fill = header_fill
                cell.font = header_font
            for entry in scenario.get("months", []):
                ws_fc.append([
                    entry["month"],
                    _fmt(entry["in"]),
                    _fmt(entry["out"]),
                    _fmt(entry["net"]),
                ])
            ws_fc.append([])

        ws_fc.column_dimensions["A"].width = 25
        for col in ["B", "C", "D"]:
            ws_fc.column_dimensions[col].width = 16

        wb.save(output_path)
        return output_path


class DashboardExporter(IReportExporter):
    """Generates normalized JSON structures for the CFO dashboard."""

    format_name: str = "dashboard_json"

    def export(
        self,
        pnl: dict[str, Any],
        cashflow: dict[str, Any],
        forecast: dict[str, Any],
        output_path: str,
        **kwargs: Any,
    ) -> str:
        """Render JSON dashboard to file."""
        data = self.build_dashboard_dict(pnl, cashflow, forecast, **kwargs)
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
        return output_path

    @staticmethod
    def build_dashboard_dict(
        pnl: dict[str, Any],
        cashflow: dict[str, Any],
        forecast: dict[str, Any],
        anomalies: list[dict[str, Any]] | None = None,
        alerts: list[dict[str, Any]] | None = None,
        narrative: str = "",
        budget: dict[str, Any] | None = None,
        multi_period: dict[str, Any] | None = None,
        tax: dict[str, Any] | None = None,
        org_id: str | None = None,
    ) -> dict[str, Any]:
        """Pure transformer creating standard dashboard dictionary."""
        base_scenario = forecast.get("scenarios", {}).get("base", {})
        growth_scenario = forecast.get("scenarios", {}).get("growth", {})
        recession_scenario = forecast.get("scenarios", {}).get("recession", {})

        return {
            "summary": {
                "revenue_cents": pnl.get("revenue", 0),
                "revenue_formatted": _fmt(pnl.get("revenue", 0)),
                "net_income_cents": pnl.get("net_income", 0),
                "net_income_formatted": _fmt(pnl.get("net_income", 0)),
                "gross_margin_pct": pnl.get("gross_margin_pct", 0.0),
                "net_margin_pct": pnl.get("net_margin_pct", 0.0),
                "operating_cashflow_cents": cashflow.get("operating", 0),
                "operating_cashflow_formatted": _fmt(cashflow.get("operating", 0)),
                "burn_rate_monthly_cents": forecast.get("burn_rate_monthly", 0),
                "runway_months": forecast.get("runway_months", "N/A"),
                "total_alerts": len(alerts or []),
                "critical_alerts": sum(1 for a in (alerts or []) if a.get("severity") == "critical"),
                "total_anomalies": len(anomalies or []),
            },
            "pnl": {
                "revenue": pnl.get("revenue", 0),
                "cogs": pnl.get("cogs", 0),
                "gross_profit": pnl.get("gross_profit", 0),
                "operating_expenses": pnl.get("operating_expenses", 0),
                "net_income": pnl.get("net_income", 0),
                "gross_margin_pct": pnl.get("gross_margin_pct", 0.0),
                "net_margin_pct": pnl.get("net_margin_pct", 0.0),
                "breakdown": pnl.get("breakdown", {}),
                "narrative": pnl.get("narrative", ""),
            },
            "cashflow": {
                "operating": cashflow.get("operating", 0),
                "investing": cashflow.get("investing", 0),
                "financing": cashflow.get("financing", 0),
                "net_change": cashflow.get("net_change", 0),
                "monthly_series": cashflow.get("monthly_series", []),
                "narrative": cashflow.get("narrative", ""),
            },
            "forecast": {
                "burn_rate_monthly": forecast.get("burn_rate_monthly", 0),
                "runway_months": forecast.get("runway_months", "N/A"),
                "scenarios": {
                    "base": {
                        "label": base_scenario.get("label", "Base Case"),
                        "series": base_scenario.get("months", []),
                    },
                    "growth": {
                        "label": growth_scenario.get("label", "Growth (+15%)"),
                        "series": growth_scenario.get("months", []),
                    },
                    "recession": {
                        "label": recession_scenario.get("label", "Recession (-20%)"),
                        "series": recession_scenario.get("months", []),
                    },
                },
                "narrative": forecast.get("narrative", ""),
            },
            "anomalies": anomalies or [],
            "alerts": alerts or [],
            "budget": budget,
            "multi_period": multi_period,
            "tax": tax,
            "executive_narrative": narrative,
            "meta": {
                "org_id": org_id,
                "version": "2.0",
            },
        }
