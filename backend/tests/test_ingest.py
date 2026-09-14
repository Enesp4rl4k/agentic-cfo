"""Files as people export them: recognised, rewritten for the parsers, and put in place.

The person who drops a staff list in does not know it is a "headcount source"
and did not name its columns `name` and `salary`. These tests use the files
such a person has: Turkish headers, ";" and Windows-1254, a bank's title rows
above the table, "45.000,00", "05.01.2024", and a TOPLAM row at the bottom.
"""
from __future__ import annotations

import ast
import csv
import io
import json
from pathlib import Path

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.services.ingest import schemas as S
from app.services.ingest.recognize import BELIRSIZ, KESIN, TANINMADI, standart_csv, tani
from app.services.ingest.table import anahtar, sayi, tarih

_APP = Path(__file__).resolve().parent.parent / "app"


def _xlsx(*rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv_rows(text: str) -> list[dict[str, str]]:
    return list(csv.DictReader(io.StringIO(text)))


BANKA = _xlsx(
    ["AKBANK T.A.Ş. HESAP HAREKETLERİ"], ["Hesap No: 1234567"], [],
    ["İşlem Tarihi", "Açıklama", "Borç", "Alacak", "Bakiye"],
    ["05.01.2024", "Ofis kirası", "28.500,00", None, "100.000,00"],
    ["06.01.2024", "EFT gelen - ABC Ltd", None, "10.000,00", "110.000,00"],
    ["TOPLAM", None, "28.500,00", "10.000,00", None],
)
PERSONEL = ("Ad Soyad;Departman;Unvan;Brüt Maaş (TL);İşe Giriş Tarihi;Durum\n"
            "Ayşe Yılmaz;Muhasebe;Uzman;45.000,00;01.03.2021;Aktif\n"
            "Mehmet Öz;Satış;Temsilci;38.250,50;15.09.2023;Aktif\n").encode("cp1254")


# ── Values ──────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("yazilan,beklenen", [
    ("45.000,00", 45000.0), ("1,234.50", 1234.5), ("-1.500,00", -1500.0), ("(1.500)", -1500.0),
    ("1.500-", -1500.0), ("12,5", 12.5), ("%12", 12.0), ("45.000", 45000.0), ("1.5", 1.5),
    (38250.5, 38250.5), ("", None), ("yok", None),
])
def test_numbers_as_written(yazilan, beklenen):
    assert sayi(yazilan) == beklenen


@pytest.mark.parametrize("yazilan,beklenen", [
    ("05.01.2024", "2024-01-05"), ("05/01/2024", "2024-01-05"), ("2024-01-05", "2024-01-05"),
    ("05.01.2024 14:30", "2024-01-05T14:30:00"), ("dün", None),
])
def test_dates_day_first(yazilan, beklenen):
    assert tarih(yazilan) == beklenen


def test_headers_fold_units_and_punctuation():
    assert anahtar("Brüt Maaş (TL)") == anahtar("brut_maas") == "brut maas"
    assert anahtar("İŞLEM TARİHİ") == "islem tarihi"


# ── Recognition ─────────────────────────────────────────────────────────────

def test_a_bank_export_under_its_title_rows():
    t = tani(BANKA, "ekstre.xlsx")
    assert (t.durum, t.tur, t.alan, t.satir_sayisi) == (KESIN, S.BANKA_EKSTRESI, "cfo", 2)
    rows = _csv_rows(standart_csv(t.tablo, t.tur)[0])
    # Debit leaves the account, credit arrives; the TOPLAM row is not a transaction.
    assert [(r["tarih"], r["tutar"]) for r in rows] == [("2024-01-05", "-28500"), ("2024-01-06", "10000")]


def test_a_turkish_staff_list_in_windows_1254_with_semicolons():
    t = tani(PERSONEL, "personel.csv")
    assert (t.durum, t.tur) == (KESIN, "headcount")
    rows = _csv_rows(standart_csv(t.tablo, t.tur)[0])
    # "Brüt Maaş" is monthly in Turkey; the parser sums `salary` as annual pay.
    assert rows[0] == {"name": "Ayşe Yılmaz", "department": "Muhasebe", "role": "Uzman",
                       "salary": "540000", "start_date": "2021-03-01", "status": "active"}
    assert rows[1]["salary"] == "459006"
    assert "×12" in t.ozet and "Brüt Maaş (TL)" in t.ozet


def test_an_annual_salary_header_is_left_as_a_year():
    t = tani("Ad Soyad,Departman,Yıllık Brüt Maaş\nAli,Satış,480.000\n".encode(), "yillik.csv")
    assert t.tur == "headcount" and "×12" not in t.ozet
    assert _csv_rows(standart_csv(t.tablo, t.tur)[0])[0]["salary"] == "480000"


def test_the_type_that_explains_more_columns_wins():
    # Name + department + salary + bonus is a pay list, not just a staff list.
    t = tani(b"Ad Soyad,Departman,Maas,Prim\nAli,Satis,40000,5000\n", "liste.csv")
    assert (t.durum, t.tur) == (KESIN, "compensation")


def test_equal_fits_are_asked_not_guessed():
    # Title + severity fits both an audit finding and a policy... only when
    # neither has more; build the tie from two schemas' required columns alone.
    t = tani(b"Politika,Durum,Bulgu,Onem\nKVKK,aktif,x,yuksek\n", "karisik.csv")
    assert t.durum == BELIRSIZ
    assert {a.tur for a in t.adaylar} == {"policies", "findings"}
    assert "seçin" in t.ozet


def test_an_unknown_table_names_what_is_missing():
    t = tani("Ad Soyad,Unvan,Telefon\nAli,Uzman,555\n".encode(), "rehber.csv")
    assert t.durum == TANINMADI
    assert "department" in t.ozet


def test_no_known_header_at_all():
    t = tani(b"foo,bar\n1,2\n", "x.csv")
    assert t.durum == TANINMADI and "başlık" in t.ozet


def test_old_excel_is_explained_in_plain_words():
    t = tani(b"\xd0\xcf\x11\xe0 old", "eski.xls")
    assert t.durum == TANINMADI and ".xlsx" in t.ozet


def test_a_workbook_uses_the_sheet_that_is_a_table():
    wb = Workbook()
    wb.active.append(["Bu dosya 2024 kampanya raporudur"])
    ws = wb.create_sheet("Veri")
    ws.append(["Kampanya adı", "Harcanan tutar", "Tıklamalar"])
    ws.append(["Yaz", 12500.5, 300])
    buf = io.BytesIO()
    wb.save(buf)
    t = tani(buf.getvalue(), "rapor.xlsx")
    assert (t.durum, t.tur, t.sayfa) == (KESIN, "campaign", "Veri")


def test_pdf_goes_to_the_financial_pipeline_and_git_log_to_cto():
    assert tani(b"%PDF-1.7", "fatura.pdf").alan == "cfo"
    log = b"commit 0123abc\nAuthor: Dev <d@x>\nDate: Mon\n\n    fix\n"
    assert (tani(log, "gecmis.txt").tur, tani(log, "gecmis.txt").alan) == ("git_log", "cto")


# ── The schemas speak the parsers' language ────────────────────────────────

_PARSER = {
    "cloud_billing": "cto/infra_agent._parse_billing_csv", "incident_log": "cto/incident_agent._parse_incident_csv",
    "sprint_data": "cto/velocity_agent._parse_sprint_csv", "headcount": "chro/headcount_agent._parse_headcount_csv",
    "attrition": "chro/attrition_agent._parse_attrition_csv",
    "compensation": "chro/compensation_agent._parse_compensation_csv",
    "campaign": "cmo/campaign_agent._parse_campaign_csv", "funnel": "cmo/funnel_agent._parse_funnel_csv",
    "cohort": "cmo/cohort_agent._parse_cohort_csv", "sla": "coo/sla_agent._parse_sla_csv",
    "process": "coo/process_agent._parse_process_csv", "resource": "coo/resource_agent._parse_resource_csv",
    "risk_register": "risk/register_agent._parse_register_csv", "loss_events": "risk/loss_agent._parse_loss_csv",
    "kri": "risk/kri_agent._parse_kri_csv", "findings": "audit/findings_agent._parse_findings_csv",
    "controls": "audit/controls_agent._parse_controls_csv", "coverage": "audit/coverage_agent._parse_coverage_csv",
    "policies": "compliance/policies_agent._parse_policies_csv",
    "violations": "compliance/violations_agent._parse_violations_csv",
    "regulations": "compliance/regulations_agent._parse_regulations_csv",
}


def _parser_columns(ref: str) -> list[list[str]]:
    mod, fn = ref.split(".")
    tree = ast.parse((_APP / "agents" / f"{mod}.py").read_text(encoding="utf-8"))
    func = next(n for n in ast.walk(tree) if isinstance(n, ast.FunctionDef) and n.name == fn)
    return [[a.value for a in c.args if isinstance(a, ast.Constant)]
            for c in ast.walk(func)
            if isinstance(c, ast.Call) and isinstance(c.func, ast.Name) and c.func.id == "_col" and c.args]


@pytest.mark.parametrize("tur", sorted(_PARSER))
def test_every_schema_column_is_one_its_parser_reads(tur):
    """A normalised column the parser does not look for is data silently dropped."""
    cols = _parser_columns(_PARSER[tur])
    sema = S.SEMA_BY_TUR[tur]
    for alan in sema.alanlar:
        assert any(alan.ad in c for c in cols), f"{tur}.{alan.ad} is not read by {_PARSER[tur]}"
    for c in cols:
        assert any(a.ad in c for a in sema.alanlar), f"{_PARSER[tur]} reads {c} but {tur} never produces it"


def test_every_parser_type_has_a_schema():
    from app.models.data_source import DOMAIN_SOURCE_KWARGS

    tabular = {t for (_d, t), kw in DOMAIN_SOURCE_KWARGS.items() if kw.endswith("_csv")}
    assert tabular == set(_PARSER)


def test_no_header_means_two_columns_in_one_schema():
    for sema in S.SEMALAR:
        seen: dict[str, str] = {}
        for alan in sema.alanlar:
            for e in alan.esler:
                k = anahtar(e)
                assert seen.setdefault(k, alan.ad) == alan.ad, f"{sema.tur}: '{e}' names two columns"


# ── Over HTTP, into a real analysis ────────────────────────────────────────

@pytest_asyncio.fixture
async def client(tmp_path, monkeypatch):
    from app.config import get_settings
    from app.database import Base, get_db
    from app.main import app

    settings = get_settings()
    monkeypatch.setattr(settings, "storage_local_path", str(tmp_path))
    monkeypatch.setattr(settings, "auto_enqueue_analysis_on_upload", False)
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    maker = async_sessionmaker(engine, expire_on_commit=False)

    async def override():
        async with maker() as s:
            yield s

    app.dependency_overrides[get_db] = override
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        c._maker = maker  # type: ignore[attr-defined]
        yield c
    app.dependency_overrides.clear()
    await engine.dispose()


async def _login(client, email="kobi@example.com"):
    from app.models.organization import Organization
    from app.models.user import User

    pw = "StrongPassword123!"
    await client.post("/api/v1/auth/register", json={"email": email, "password": pw, "full_name": "K"})
    async with client._maker() as db:
        user = (await db.execute(select(User).where(User.email == email))).scalar_one()
        org = Organization(name=email, slug=email.split("@")[0], regional_packs=["tr"])
        db.add(org)
        await db.flush()
        user.org_id, user.role = org.id, "owner"
        await db.commit()
    token = (await client.post("/api/v1/auth/login", json={"email": email, "password": pw})).json()["data"]["access_token"]
    return {"Authorization": f"Bearer {token}"}


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.mark.asyncio
async def test_drop_a_statement_and_three_hr_files_together_and_the_hr_page_analyses_them(client):
    h = await _login(client)
    ayrilan = ("Ad Soyad;Departman;Ayrılış Tarihi;Ayrılış Nedeni\nCan Ak;Satış;30.06.2024;Maaş\n").encode("cp1254")
    ucret = _xlsx(["Ad Soyad", "Departman", "Brüt Maaş", "Prim"], ["Ayşe Yılmaz", "Muhasebe", 45000, 4500],
                  ["Mehmet Öz", "Satış", 38250.5, 3000])
    files = [
        ("files", ("personel.csv", PERSONEL, "text/csv")),
        ("files", ("ayrilanlar.csv", ayrilan, "text/csv")),
        ("files", ("ucretler.xlsx", ucret, XLSX)),
        ("files", ("ekstre.xlsx", BANKA, XLSX)),
    ]
    r = await client.post("/api/v1/veri/ekle", files=files, headers=h)
    assert r.status_code == 201, r.text
    body = r.json()["data"]
    by = {d["dosya"]: d for d in body["dosyalar"]}
    assert {k: v["durum"] for k, v in by.items()} == {k: "eklendi" for k in by}
    assert by["ekstre.xlsx"]["job_id"] == body["job_id"]
    assert {by[f]["tanima"]["tur"] for f in ("personel.csv", "ayrilanlar.csv", "ucretler.xlsx")} == {
        "headcount", "attrition", "compensation"}
    assert by["personel.csv"]["sayfa"] == "/chro"

    # The statement was stored as its parser reads it, and the parser reads it.
    from app.agents.data_ingestion import run_data_ingestion
    from app.agents.orchestration import domain_analysis as da
    from app.agents.state import AgentRunConfig
    from app.models.analysis_job import AnalysisJob

    async with client._maker() as db:
        job = await db.get(AnalysisJob, body["job_id"])
        ingested = await run_data_ingestion({"file_path": job.file_path, "file_type": job.file_type,
                                             "job_id": job.id}, AgentRunConfig())
        assert [(t["type"], t["amount_cents"]) for t in ingested.patch["transactions"]] == [
            ("expense", 2_850_000), ("income", 1_000_000)]

        out = await da.analiz_et(db, job, "chro")
    assert out["durum"] == da.ANALIZ_EDILDI, out.get("neden")
    headcount = out["sonuc"]["headcount"]
    assert headcount["total_headcount"] == 2
    # Turkish monthly salaries, read as numbers and as a year: (45,000 + 38,250.50) × 12, in kuruş.
    assert headcount["total_annual_payroll"] == 99_900_600


@pytest.mark.asyncio
async def test_an_ambiguous_file_is_not_saved_until_the_person_chooses(client):
    h = await _login(client)
    await client.post("/api/v1/veri/ekle", files=[("files", ("ekstre.xlsx", BANKA, XLSX))], headers=h)
    karisik = b"Politika,Durum,Bulgu,Onem\nKVKK,aktif,x,yuksek\n"

    r = await client.post("/api/v1/veri/ekle", files=[("files", ("k.csv", karisik, "text/csv"))], headers=h)
    d = r.json()["data"]["dosyalar"][0]
    assert d["durum"] == "secim_gerekli" and "source_id" not in d

    r = await client.post("/api/v1/veri/ekle", files=[("files", ("k.csv", karisik, "text/csv"))],
                          data={"secimler": json.dumps({"k.csv": "policies"})}, headers=h)
    d = r.json()["data"]["dosyalar"][0]
    assert (d["durum"], d["sayfa"]) == ("eklendi", "/compliance")


@pytest.mark.asyncio
async def test_a_domain_file_without_any_analysis_says_what_to_add_first(client):
    h = await _login(client)
    r = await client.post("/api/v1/veri/ekle", files=[("files", ("personel.csv", PERSONEL, "text/csv"))], headers=h)
    d = r.json()["data"]["dosyalar"][0]
    assert d["durum"] == "finansal_dosya_gerekli" and "banka ekstresi" in d["mesaj"]


@pytest.mark.asyncio
async def test_another_organisations_job_is_refused_before_anything_is_saved(client):
    owner = await _login(client, "a@example.com")
    job_id = (await client.post("/api/v1/veri/ekle", files=[("files", ("e.xlsx", BANKA, XLSX))],
                                headers=owner)).json()["data"]["job_id"]
    stranger = await _login(client, "b@example.com")
    r = await client.post("/api/v1/veri/ekle", files=[("files", ("personel.csv", PERSONEL, "text/csv"))],
                          data={"job_id": job_id}, headers=stranger)
    assert r.status_code == 404
    from app.models.data_source import DataSource

    async with client._maker() as db:
        assert (await db.execute(select(DataSource))).scalars().all() == []


@pytest.mark.asyncio
async def test_unsupported_and_oversized_files_are_refused_in_words(client):
    h = await _login(client)
    r = await client.post("/api/v1/veri/ekle", files=[("files", ("virus.exe", b"MZ", "application/octet-stream"))],
                          headers=h)
    d = r.json()["data"]["dosyalar"][0]
    assert d["durum"] == "reddedildi" and "Excel" in d["mesaj"]
